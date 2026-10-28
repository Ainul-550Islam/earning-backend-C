"""
Gamification Data Export — Comprehensive data export functionality.

This module provides data export capabilities for the gamification system,
including various export formats, filtering options, and compliance features.
"""

from __future__ import annotations

import csv
import json
import logging
import zipfile
import io
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple, Union

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.db import models
from django.db.models import Q, F, Func, Value, Case, When, Sum, Count, Avg, Max, Min
from django.http import HttpResponse
from django.utils import timezone

from api.gamification.models import (
    Badge, UserBadge, DailyChallenge, UserDailyChallenge,
    Streak, LevelConfig, ContestCycle, UserAchievement,
    WebhookEndpoint, WebhookEvent, WebhookDelivery
)
from api.gamification.choices import (
    BadgeRarity, BadgeCriteriaType, ChallengeAction,
    ContestCycleStatus, AchievementType, WebhookStatus, WebhookEventType
)
from api.gamification.exceptions import GamificationServiceError
from api.gamification.repositories import (
    BadgeRepository, UserBadgeRepository, StreakRepository,
    UserAchievementRepository, WebhookDeliveryRepository
)
from api.gamification.encryption import data_masking, compliance_checker

User = get_user_model()
logger = logging.getLogger(__name__)


class ExportFormat:
    """Supported export formats."""
    CSV = "csv"
    JSON = "json"
    EXCEL = "excel"
    PDF = "pdf"


class DataExporter:
    """Base class for data exporters."""
    
    def __init__(self, export_format: str = ExportFormat.CSV):
        self.export_format = export_format
        self.logger = logging.getLogger(__name__)
    
    def export(self, data: List[Dict[str, Any]], filename: str,
                include_headers: bool = True, **kwargs) -> Union[HttpResponse, bytes]:
        """Export data to specified format."""
        if self.export_format == ExportFormat.CSV:
            return self._export_csv(data, filename, include_headers)
        elif self.export_format == ExportFormat.JSON:
            return self._export_json(data, filename)
        elif self.export_format == ExportFormat.EXCEL:
            return self._export_excel(data, filename, include_headers)
        elif self.export_format == ExportFormat.PDF:
            return self._export_pdf(data, filename, **kwargs)
        else:
            raise ValueError(f"Unsupported export format: {self.export_format}")
    
    def _export_csv(self, data: List[Dict[str, Any]], filename: str,
                   include_headers: bool = True) -> HttpResponse:
        """Export data to CSV format."""
        if not data:
            raise GamificationServiceError("No data to export")
        
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
        
        writer = csv.writer(response)
        
        # Write headers
        if include_headers:
            headers = list(data[0].keys())
            writer.writerow(headers)
        
        # Write data rows
        for row in data:
            writer.writerow([str(row.get(header, '')) for header in headers])
        
        return response
    
    def _export_json(self, data: List[Dict[str, Any]], filename: str) -> HttpResponse:
        """Export data to JSON format."""
        response = HttpResponse(content_type='application/json')
        response['Content-Disposition'] = f'attachment; filename="{filename}.json"'
        
        json_data = {
            'export_info': {
                'filename': filename,
                'format': 'json',
                'exported_at': timezone.now().isoformat(),
                'record_count': len(data)
            },
            'data': data
        }
        
        response.write(json.dumps(json_data, indent=2, default=str))
        
        return response
    
    def _export_excel(self, data: List[Dict[str, Any]], filename: str,
                     include_headers: bool = True) -> HttpResponse:
        """Export data to Excel format."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise GamificationServiceError("openpyxl package is required for Excel export")
        
        if not data:
            raise GamificationServiceError("No data to export")
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Gamification Data"
        
        # Write headers
        if include_headers:
            headers = list(data[0].keys())
            header_cell = ws.cell(row=1, column=1)
            header_cell.value = "Gamification Data Export"
            header_cell.font = Font(bold=True, size=14)
            
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=2, column=col)
                cell.value = header
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
            
            # Write data rows
            for row_idx, row_data in enumerate(data, 3):
                for col, header in enumerate(headers, 1):
                    ws.cell(row=row_idx, column=col, value=str(row_data.get(header, '')))
        else:
            # Write data rows without headers
            for row_idx, row_data in enumerate(data, 1):
                for col, value in enumerate(row_data.values(), 1):
                    ws.cell(row=row_idx, column=col, value=str(value))
        
        # Save to response
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        
        wb.save(response)
        
        return response
    
    def _export_pdf(self, data: List[Dict[str, Any]], filename: str,
                   **kwargs) -> HttpResponse:
        """Export data to PDF format."""
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter, landscape
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
        except ImportError:
            raise GamificationServiceError("reportlab package is required for PDF export")
        
        if not data:
            raise GamificationServiceError("No data to export")
        
        # Create PDF buffer
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter))
        
        # Get styles
        styles = getSampleStyleSheet()
        elements = []
        
        # Add title
        title = Paragraph(f"Gamification Data Export - {filename}", styles['Title'])
        elements.append(title)
        elements.append(Spacer(1, 12))
        
        # Prepare table data
        headers = list(data[0].keys())
        table_data = [headers]
        
        for row in data:
            table_data.append([str(row.get(header, '')) for header in headers])
        
        # Create table
        table = Table(table_data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 14),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        elements.append(table)
        
        # Build PDF
        doc.build(elements)
        
        # Create response
        buffer.seek(0)
        response = HttpResponse(buffer.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
        
        return response


class GamificationDataExporter:
    """Specialized exporter for gamification data."""
    
    def __init__(self, export_format: str = ExportFormat.CSV):
        self.exporter = DataExporter(export_format)
        self.logger = logging.getLogger(__name__)
    
    def export_user_badges(self, user_id: Optional[str] = None,
                          date_range: Optional[Tuple[datetime, datetime]] = None,
                          include_sensitive: bool = False) -> HttpResponse:
        """Export user badges data."""
        try:
            # Get data
            queryset = UserBadge.objects.select_related('badge', 'user')
            
            if user_id:
                queryset = queryset.filter(user_id=user_id)
            
            if date_range:
                start_date, end_date = date_range
                queryset = queryset.filter(earned_at__range=[start_date, end_date])
            
            # Convert to list of dictionaries
            data = []
            for user_badge in queryset:
                row = {
                    'id': str(user_badge.id),
                    'user_id': str(user_badge.user.id),
                    'username': user_badge.user.username,
                    'badge_id': str(user_badge.badge.id),
                    'badge_name': user_badge.badge.name,
                    'badge_description': user_badge.badge.description,
                    'badge_rarity': user_badge.badge.rarity,
                    'badge_criteria_type': user_badge.badge.criteria_type,
                    'points_bonus': user_badge.badge.points_bonus,
                    'is_pinned': user_badge.is_pinned,
                    'earned_at': user_badge.earned_at.isoformat(),
                    'metadata': json.dumps(user_badge.metadata or {})
                }
                
                # Apply data masking if not including sensitive data
                if not include_sensitive:
                    row['username'] = data_masking.mask_general(row['username'], 4)
                
                data.append(row)
            
            # Export
            filename = f"user_badges_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
            return self.exporter.export(data, filename)
            
        except Exception as e:
            self.logger.error(f"Error exporting user badges: {e}")
            raise GamificationServiceError(f"Failed to export user badges: {e}")
    
    def export_user_streaks(self, user_id: Optional[str] = None,
                           date_range: Optional[Tuple[datetime, datetime]] = None,
                           include_sensitive: bool = False) -> HttpResponse:
        """Export user streaks data."""
        try:
            # Get data
            queryset = Streak.objects.select_related('user')
            
            if user_id:
                queryset = queryset.filter(user_id=user_id)
            
            if date_range:
                start_date, end_date = date_range
                queryset = queryset.filter(last_activity_date__range=[start_date, end_date])
            
            # Convert to list of dictionaries
            data = []
            for streak in queryset:
                row = {
                    'id': str(streak.id),
                    'user_id': str(streak.user.id),
                    'username': streak.user.username,
                    'current_days': streak.current_days,
                    'longest_days': streak.longest_days,
                    'last_activity_date': streak.last_activity_date.isoformat(),
                    'is_active': streak.is_active,
                    'created_at': streak.created_at.isoformat(),
                    'updated_at': streak.updated_at.isoformat()
                }
                
                # Apply data masking if not including sensitive data
                if not include_sensitive:
                    row['username'] = data_masking.mask_general(row['username'], 4)
                
                data.append(row)
            
            # Export
            filename = f"user_streaks_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
            return self.exporter.export(data, filename)
            
        except Exception as e:
            self.logger.error(f"Error exporting user streaks: {e}")
            raise GamificationServiceError(f"Failed to export user streaks: {e}")
    
    def export_user_challenges(self, user_id: Optional[str] = None,
                             date_range: Optional[Tuple[datetime, datetime]] = None,
                             include_sensitive: bool = False) -> HttpResponse:
        """Export user challenges data."""
        try:
            # Get data
            queryset = UserDailyChallenge.objects.select_related('user', 'challenge')
            
            if user_id:
                queryset = queryset.filter(user_id=user_id)
            
            if date_range:
                start_date, end_date = date_range
                queryset = queryset.filter(
                    Q(challenge__date__range=[start_date, end_date]) |
                    Q(completed_at__range=[start_date, end_date])
                )
            
            # Convert to list of dictionaries
            data = []
            for user_challenge in queryset:
                row = {
                    'id': str(user_challenge.id),
                    'user_id': str(user_challenge.user.id),
                    'username': user_challenge.user.username,
                    'challenge_id': str(user_challenge.challenge.id),
                    'challenge_title': user_challenge.challenge.title,
                    'challenge_description': user_challenge.challenge.description,
                    'target_action': user_challenge.challenge.target_action,
                    'target_value': user_challenge.challenge.target_value,
                    'bonus_points': user_challenge.challenge.bonus_points,
                    'current_progress': user_challenge.current_progress,
                    'is_completed': user_challenge.is_completed,
                    'completed_at': user_challenge.completed_at.isoformat() if user_challenge.completed_at else None,
                    'challenge_date': user_challenge.challenge.date.isoformat(),
                    'created_at': user_challenge.created_at.isoformat(),
                    'updated_at': user_challenge.updated_at.isoformat()
                }
                
                # Apply data masking if not including sensitive data
                if not include_sensitive:
                    row['username'] = data_masking.mask_general(row['username'], 4)
                
                data.append(row)
            
            # Export
            filename = f"user_challenges_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
            return self.exporter.export(data, filename)
            
        except Exception as e:
            self.logger.error(f"Error exporting user challenges: {e}")
            raise GamificationServiceError(f"Failed to export user challenges: {e}")
    
    def export_user_achievements(self, user_id: Optional[str] = None,
                                date_range: Optional[Tuple[datetime, datetime]] = None,
                                include_sensitive: bool = False) -> HttpResponse:
        """Export user achievements data."""
        try:
            # Get data
            queryset = UserAchievement.objects.select_related('user', 'contest_cycle')
            
            if user_id:
                queryset = queryset.filter(user_id=user_id)
            
            if date_range:
                start_date, end_date = date_range
                queryset = queryset.filter(awarded_at__range=[start_date, end_date])
            
            # Convert to list of dictionaries
            data = []
            for achievement in queryset:
                row = {
                    'id': str(achievement.id),
                    'user_id': str(achievement.user.id),
                    'username': achievement.user.username,
                    'achievement_type': achievement.achievement_type,
                    'points_awarded': achievement.points_awarded,
                    'rank': achievement.rank,
                    'is_awarded': achievement.is_awarded,
                    'awarded_at': achievement.awarded_at.isoformat() if achievement.awarded_at else None,
                    'contest_cycle_id': str(achievement.contest_cycle.id) if achievement.contest_cycle else None,
                    'contest_name': achievement.contest_cycle.name if achievement.contest_cycle else None,
                    'metadata': json.dumps(achievement.metadata or {}),
                    'created_at': achievement.created_at.isoformat(),
                    'updated_at': achievement.updated_at.isoformat()
                }
                
                # Apply data masking if not including sensitive data
                if not include_sensitive:
                    row['username'] = data_masking.mask_general(row['username'], 4)
                
                data.append(row)
            
            # Export
            filename = f"user_achievements_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
            return self.exporter.export(data, filename)
            
        except Exception as e:
            self.logger.error(f"Error exporting user achievements: {e}")
            raise GamificationServiceError(f"Failed to export user achievements: {e}")
    
    def export_leaderboard(self, scope: str = "global", limit: int = 100,
                          date_range: Optional[Tuple[datetime, datetime]] = None,
                          include_sensitive: bool = False) -> HttpResponse:
        """Export leaderboard data."""
        try:
            # Get leaderboard data
            from api.gamification.helpers import LeaderboardHelper
            leaderboard = LeaderboardHelper.get_leaderboard(scope, limit)
            
            # Convert to list of dictionaries
            data = []
            for entry in leaderboard:
                row = {
                    'rank': entry['rank'],
                    'user_id': entry['user_id'],
                    'username': entry['username'],
                    'total_points': entry['total_points'],
                    'badge_count': entry['badge_count'],
                    'streak_days': entry['streak_days'],
                    'level': entry['level'],
                    'scope': scope,
                    'generated_at': timezone.now().isoformat()
                }
                
                # Apply data masking if not including sensitive data
                if not include_sensitive:
                    row['username'] = data_masking.mask_general(row['username'], 4)
                
                data.append(row)
            
            # Export
            filename = f"leaderboard_{scope}_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
            return self.exporter.export(data, filename)
            
        except Exception as e:
            self.logger.error(f"Error exporting leaderboard: {e}")
            raise GamificationServiceError(f"Failed to export leaderboard: {e}")
    
    def export_system_stats(self, date_range: Optional[Tuple[datetime, datetime]] = None) -> HttpResponse:
        """Export system statistics."""
        try:
            # Get system statistics
            from api.gamification.analytics import system_analytics
            stats = system_analytics.get_system_overview()
            
            # Flatten nested data for export
            data = []
            
            # User statistics
            for key, value in stats['user_statistics'].items():
                data.append({
                    'category': 'user_statistics',
                    'metric': key,
                    'value': value,
                    'generated_at': stats['generated_at']
                })
            
            # Badge statistics
            for key, value in stats['badge_statistics'].items():
                data.append({
                    'category': 'badge_statistics',
                    'metric': key,
                    'value': value,
                    'generated_at': stats['generated_at']
                })
            
            # Challenge statistics
            for key, value in stats['challenge_statistics'].items():
                data.append({
                    'category': 'challenge_statistics',
                    'metric': key,
                    'value': value,
                    'generated_at': stats['generated_at']
                })
            
            # Streak statistics
            for key, value in stats['streak_statistics'].items():
                data.append({
                    'category': 'streak_statistics',
                    'metric': key,
                    'value': value,
                    'generated_at': stats['generated_at']
                })
            
            # Achievement statistics
            for key, value in stats['achievement_statistics'].items():
                data.append({
                    'category': 'achievement_statistics',
                    'metric': key,
                    'value': value,
                    'generated_at': stats['generated_at']
                })
            
            # Export
            filename = f"system_stats_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
            return self.exporter.export(data, filename)
            
        except Exception as e:
            self.logger.error(f"Error exporting system stats: {e}")
            raise GamificationServiceError(f"Failed to export system stats: {e}")
    
    def export_analytics_data(self, analytics_type: str = "engagement",
                             user_id: Optional[str] = None,
                             date_range: Optional[Tuple[datetime, datetime]] = None,
                             include_sensitive: bool = False) -> HttpResponse:
        """Export analytics data."""
        try:
            from api.gamification.analytics import UserAnalytics, system_analytics
            
            if user_id:
                # Get user analytics
                user_analytics = UserAnalytics(user_id)
                
                if analytics_type == "engagement":
                    data = [user_analytics.get_user_engagement_summary()]
                elif analytics_type == "progression":
                    data = [user_analytics.get_user_progression_analysis()]
                elif analytics_type == "behavior":
                    data = [user_analytics.get_user_behavior_patterns()]
                else:
                    raise ValueError(f"Unknown analytics type: {analytics_type}")
                
                # Flatten nested data
                flattened_data = []
                for item in data:
                    for key, value in item.items():
                        if isinstance(value, (dict, list)):
                            flattened_data.append({
                                'user_id': user_id,
                                'category': key,
                                'data': json.dumps(value) if isinstance(value, (dict, list)) else str(value),
                                'generated_at': item.get('generated_at', timezone.now().isoformat())
                            })
                        else:
                            flattened_data.append({
                                'user_id': user_id,
                                'category': analytics_type,
                                'metric': key,
                                'value': value,
                                'generated_at': item.get('generated_at', timezone.now().isoformat())
                            })
                
                data = flattened_data
            else:
                # Get system analytics
                if analytics_type == "engagement":
                    data = system_analytics.get_engagement_analytics()
                elif analytics_type == "performance":
                    data = system_analytics.get_performance_analytics()
                elif analytics_type == "retention":
                    data = system_analytics.get_retention_analytics()
                else:
                    raise ValueError(f"Unknown analytics type: {analytics_type}")
                
                # Flatten nested data
                flattened_data = []
                for key, value in data.items():
                    if isinstance(value, (dict, list)):
                        flattened_data.append({
                            'category': analytics_type,
                            'subcategory': key,
                            'data': json.dumps(value) if isinstance(value, (dict, list)) else str(value),
                            'generated_at': data.get('generated_at', timezone.now().isoformat())
                        })
                    else:
                        flattened_data.append({
                            'category': analytics_type,
                            'metric': key,
                            'value': value,
                            'generated_at': data.get('generated_at', timezone.now().isoformat())
                        })
                
                data = flattened_data
            
            # Apply data masking if not including sensitive data
            if not include_sensitive:
                for row in data:
                    if 'user_id' in row:
                        row['user_id'] = data_masking.mask_general(row['user_id'], 4)
            
            # Export
            filename = f"analytics_{analytics_type}_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
            return self.exporter.export(data, filename)
            
        except Exception as e:
            self.logger.error(f"Error exporting analytics data: {e}")
            raise GamificationServiceError(f"Failed to export analytics data: {e}")
    
    def export_compliance_data(self, compliance_type: str = "gdpr",
                              user_id: Optional[str] = None,
                              include_sensitive: bool = False) -> HttpResponse:
        """Export compliance data."""
        try:
            if compliance_type == "gdpr":
                # Get user data for GDPR compliance
                if not user_id:
                    raise ValueError("User ID is required for GDPR export")
                
                # Collect all user data
                user_data = {
                    'user_info': {},
                    'badges': [],
                    'streaks': [],
                    'challenges': [],
                    'achievements': [],
                    'analytics': []
                }
                
                # Get basic user info
                try:
                    user = User.objects.get(id=user_id)
                    user_data['user_info'] = {
                        'id': str(user.id),
                        'username': user.username,
                        'email': user.email if include_sensitive else data_masking.mask_email(user.email),
                        'first_name': user.first_name,
                        'last_name': user.last_name,
                        'date_joined': user.date_joined.isoformat(),
                        'last_login': user.last_login.isoformat() if user.last_login else None,
                        'is_active': user.is_active
                    }
                except User.DoesNotExist:
                    raise GamificationServiceError("User not found")
                
                # Get user badges
                user_badges = UserBadge.objects.filter(user_id=user_id).select_related('badge')
                for user_badge in user_badges:
                    user_data['badges'].append({
                        'badge_name': user_badge.badge.name,
                        'badge_description': user_badge.badge.description,
                        'badge_rarity': user_badge.badge.rarity,
                        'points_bonus': user_badge.badge.points_bonus,
                        'earned_at': user_badge.earned_at.isoformat(),
                        'is_pinned': user_badge.is_pinned
                    })
                
                # Get user streaks
                streak = Streak.objects.filter(user_id=user_id).first()
                if streak:
                    user_data['streaks'].append({
                        'current_days': streak.current_days,
                        'longest_days': streak.longest_days,
                        'last_activity_date': streak.last_activity_date.isoformat(),
                        'is_active': streak.is_active
                    })
                
                # Get user challenges
                user_challenges = UserDailyChallenge.objects.filter(user_id=user_id).select_related('challenge')
                for user_challenge in user_challenges:
                    user_data['challenges'].append({
                        'challenge_title': user_challenge.challenge.title,
                        'target_action': user_challenge.challenge.target_action,
                        'target_value': user_challenge.challenge.target_value,
                        'current_progress': user_challenge.current_progress,
                        'is_completed': user_challenge.is_completed,
                        'completed_at': user_challenge.completed_at.isoformat() if user_challenge.completed_at else None,
                        'challenge_date': user_challenge.challenge.date.isoformat()
                    })
                
                # Get user achievements
                achievements = UserAchievement.objects.filter(user_id=user_id).select_related('contest_cycle')
                for achievement in achievements:
                    user_data['achievements'].append({
                        'achievement_type': achievement.achievement_type,
                        'points_awarded': achievement.points_awarded,
                        'rank': achievement.rank,
                        'is_awarded': achievement.is_awarded,
                        'awarded_at': achievement.awarded_at.isoformat() if achievement.awarded_at else None
                    })
                
                # Get user analytics summary
                from api.gamification.analytics import UserAnalytics
                user_analytics = UserAnalytics(user_id)
                user_data['analytics'] = [
                    user_analytics.get_user_engagement_summary(),
                    user_analytics.get_user_progression_analysis(),
                    user_analytics.get_user_behavior_patterns()
                ]
                
                # Check compliance
                compliance_result = compliance_checker.check_gdpr_compliance("user_data", user_id, True)
                user_data['compliance'] = compliance_result
                
                # Convert to list format for export
                data = [{
                    'user_id': user_id,
                    'data_type': 'gdpr_export',
                    'export_data': json.dumps(user_data),
                    'compliant': compliance_result['compliant'],
                    'violations': json.dumps(compliance_result['violations']),
                    'exported_at': timezone.now().isoformat()
                }]
                
            else:
                raise ValueError(f"Unknown compliance type: {compliance_type}")
            
            # Export
            filename = f"compliance_{compliance_type}_{user_id}_{timezone.now().strftime('%Y%m%d_%H%M%S')}"
            return self.exporter.export(data, filename)
            
        except Exception as e:
            self.logger.error(f"Error exporting compliance data: {e}")
            raise GamificationServiceError(f"Failed to export compliance data: {e}")
    
    def export_bulk_data(self, export_types: List[str],
                        user_id: Optional[str] = None,
                        date_range: Optional[Tuple[datetime, datetime]] = None,
                        include_sensitive: bool = False,
                        compress: bool = True) -> HttpResponse:
        """Export multiple data types in bulk."""
        try:
            exports = {}
            
            # Export each requested type
            for export_type in export_types:
                if export_type == "badges":
                    exports["badges"] = self.export_user_badges(user_id, date_range, include_sensitive)
                elif export_type == "streaks":
                    exports["streaks"] = self.export_user_streaks(user_id, date_range, include_sensitive)
                elif export_type == "challenges":
                    exports["challenges"] = self.export_user_challenges(user_id, date_range, include_sensitive)
                elif export_type == "achievements":
                    exports["achievements"] = self.export_user_achievements(user_id, date_range, include_sensitive)
                elif export_type == "leaderboard":
                    exports["leaderboard"] = self.export_leaderboard(include_sensitive=include_sensitive)
                elif export_type == "system_stats":
                    exports["system_stats"] = self.export_system_stats(date_range)
                elif export_type == "analytics":
                    exports["analytics"] = self.export_analytics_data(include_sensitive=include_sensitive)
                elif export_type == "compliance":
                    if not user_id:
                        raise ValueError("User ID is required for compliance export")
                    exports["compliance"] = self.export_compliance_data("gdpr", user_id, include_sensitive)
                else:
                    self.logger.warning(f"Unknown export type: {export_type}")
            
            if compress and len(exports) > 1:
                # Create ZIP file with multiple exports
                return self._create_zip_export(exports, user_id)
            elif len(exports) == 1:
                # Return single export
                return list(exports.values())[0]
            else:
                raise GamificationServiceError("No valid export types specified")
                
        except Exception as e:
            self.logger.error(f"Error in bulk export: {e}")
            raise GamificationServiceError(f"Failed to perform bulk export: {e}")
    
    def _create_zip_export(self, exports: Dict[str, HttpResponse], user_id: Optional[str] = None) -> HttpResponse:
        """Create ZIP file containing multiple exports."""
        try:
            # Create ZIP buffer
            zip_buffer = io.BytesIO()
            
            with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                for export_name, export_response in exports.items():
                    # Get content from response
                    content = export_response.content
                    
                    # Get filename from content disposition
                    content_disposition = export_response.get('Content-Disposition', '')
                    if 'filename=' in content_disposition:
                        filename = content_disposition.split('filename=')[1].strip('"')
                    else:
                        filename = f"{export_name}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.csv"
                    
                    # Add to ZIP
                    zip_file.writestr(filename, content)
            
            # Create response
            zip_buffer.seek(0)
            response = HttpResponse(zip_buffer.read(), content_type='application/zip')
            
            zip_filename = f"gamification_export_{user_id or 'bulk'}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.zip"
            response['Content-Disposition'] = f'attachment; filename="{zip_filename}"'
            
            return response
            
        except Exception as e:
            self.logger.error(f"Error creating ZIP export: {e}")
            raise GamificationServiceError(f"Failed to create ZIP export: {e}")


# Convenience functions
def export_user_badges(user_id: Optional[str] = None,
                       export_format: str = ExportFormat.CSV,
                       date_range: Optional[Tuple[datetime, datetime]] = None,
                       include_sensitive: bool = False) -> HttpResponse:
    """Export user badges data."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_user_badges(user_id, date_range, include_sensitive)


def export_user_streaks(user_id: Optional[str] = None,
                        export_format: str = ExportFormat.CSV,
                        date_range: Optional[Tuple[datetime, datetime]] = None,
                        include_sensitive: bool = False) -> HttpResponse:
    """Export user streaks data."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_user_streaks(user_id, date_range, include_sensitive)


def export_user_challenges(user_id: Optional[str] = None,
                          export_format: str = ExportFormat.CSV,
                          date_range: Optional[Tuple[datetime, datetime]] = None,
                          include_sensitive: bool = False) -> HttpResponse:
    """Export user challenges data."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_user_challenges(user_id, date_range, include_sensitive)


def export_user_achievements(user_id: Optional[str] = None,
                            export_format: str = ExportFormat.CSV,
                            date_range: Optional[Tuple[datetime, datetime]] = None,
                            include_sensitive: bool = False) -> HttpResponse:
    """Export user achievements data."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_user_achievements(user_id, date_range, include_sensitive)


def export_leaderboard(scope: str = "global", limit: int = 100,
                      export_format: str = ExportFormat.CSV,
                      include_sensitive: bool = False) -> HttpResponse:
    """Export leaderboard data."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_leaderboard(scope, limit, include_sensitive=include_sensitive)


def export_system_stats(export_format: str = ExportFormat.CSV,
                       date_range: Optional[Tuple[datetime, datetime]] = None) -> HttpResponse:
    """Export system statistics."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_system_stats(date_range)


def export_analytics_data(analytics_type: str = "engagement",
                          user_id: Optional[str] = None,
                          export_format: str = ExportFormat.CSV,
                          date_range: Optional[Tuple[datetime, datetime]] = None,
                          include_sensitive: bool = False) -> HttpResponse:
    """Export analytics data."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_analytics_data(analytics_type, user_id, date_range, include_sensitive)


def export_gdpr_data(user_id: str, export_format: str = ExportFormat.JSON,
                     include_sensitive: bool = False) -> HttpResponse:
    """Export GDPR compliance data."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_compliance_data("gdpr", user_id, include_sensitive)


def export_bulk_data(export_types: List[str],
                    user_id: Optional[str] = None,
                    export_format: str = ExportFormat.CSV,
                    date_range: Optional[Tuple[datetime, datetime]] = None,
                    include_sensitive: bool = False,
                    compress: bool = True) -> HttpResponse:
    """Export multiple data types in bulk."""
    exporter = GamificationDataExporter(export_format)
    return exporter.export_bulk_data(export_types, user_id, date_range, include_sensitive, compress)
