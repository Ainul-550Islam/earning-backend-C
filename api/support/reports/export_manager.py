"""
api/support/reports/export_manager.py
FILE 53/70 — Universal Export Manager (CSV/Excel)

All reports → CSV export।
Excel support: pip install openpyxl
"""
import logging
import csv
from io import StringIO, BytesIO
from django.utils import timezone
logger = logging.getLogger(__name__)

class ExportManager:
    def export_tickets_csv(self, days: int = 30, filters: dict = None) -> tuple:
        from api.support.reports.ticket_report import ticket_report
        return ticket_report.export_csv(days, filters)

    def export_agents_csv(self, days: int = 30) -> tuple:
        from api.support.reports.agent_report import agent_report
        return agent_report.export_csv(days)

    def export_csat_csv(self, days: int = 30) -> tuple:
        from api.support.models import TicketRating
        from django.db.models import Avg
        since   = timezone.now() - __import__('datetime', fromlist=['timedelta']).timedelta(days=days)
        ratings = TicketRating.objects.filter(created_at__gte=since).select_related('ticket','user')
        output  = StringIO()
        writer  = csv.writer(output)
        writer.writerow(['Ticket ID','User','Rating','Comment','Agent','Created At'])
        for r in ratings:
            writer.writerow([
                r.ticket.ticket_id, r.user.username, r.rating, r.comment,
                r.ticket.assigned_to.username if r.ticket.assigned_to else '',
                r.created_at.strftime('%Y-%m-%d %H:%M'),
            ])
        filename = f'csat_report_{timezone.now().strftime("%Y%m%d")}.csv'
        return filename, output.getvalue().encode('utf-8-sig')

    def export_sla_csv(self, days: int = 30) -> tuple:
        from api.support.models import SLABreach
        from datetime import timedelta
        since   = timezone.now() - timedelta(days=days)
        breaches= SLABreach.objects.filter(breached_at__gte=since).select_related('ticket')
        output  = StringIO()
        writer  = csv.writer(output)
        writer.writerow(['Ticket ID','Priority','Breach Type','Due At','Breached At','Overdue Hours','Resolved'])
        for b in breaches:
            writer.writerow([
                b.ticket.ticket_id, b.ticket.priority, b.breach_type,
                b.due_at.strftime('%Y-%m-%d %H:%M'),
                b.breached_at.strftime('%Y-%m-%d %H:%M'),
                round(b.overdue_hours, 1), 'Yes' if b.resolved else 'No',
            ])
        filename = f'sla_report_{timezone.now().strftime("%Y%m%d")}.csv'
        return filename, output.getvalue().encode('utf-8-sig')

    def export_excel(self, report_type: str, days: int = 30) -> tuple:
        """Excel export — openpyxl দিয়ে"""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = openpyxl.Workbook()
            ws = wb.active

            if report_type == 'tickets':
                filename, csv_bytes = self.export_tickets_csv(days)
                ws.title = 'Tickets'
            elif report_type == 'agents':
                filename, csv_bytes = self.export_agents_csv(days)
                ws.title = 'Agents'
            elif report_type == 'csat':
                filename, csv_bytes = self.export_csat_csv(days)
                ws.title = 'CSAT'
            else:
                return None, None

            # Parse CSV and write to Excel
            import csv as csv_mod
            from io import StringIO
            reader = csv_mod.reader(StringIO(csv_bytes.decode('utf-8-sig')))
            for i, row in enumerate(reader, 1):
                for j, val in enumerate(row, 1):
                    cell = ws.cell(row=i, column=j, value=val)
                    if i == 1:
                        cell.font       = Font(bold=True, color='FFFFFF')
                        cell.fill       = PatternFill('solid', fgColor='6366F1')
                        cell.alignment  = Alignment(horizontal='center')

            # Auto-width
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

            output   = BytesIO()
            wb.save(output)
            xlsx_filename = filename.replace('.csv', '.xlsx')
            return xlsx_filename, output.getvalue()

        except ImportError:
            logger.warning('openpyxl not installed. pip install openpyxl')
            return None, None

export_manager = ExportManager()