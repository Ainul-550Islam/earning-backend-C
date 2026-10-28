"""
Common Imports for Security Module
Consolidated imports to reduce duplication across security files
"""

# Standard library imports
import hashlib
from datetime import datetime, timedelta, date
import logging
import json
import uuid
import secrets
import string
import re
import ipaddress
import base64
import hmac
import time
import datetime
import decimal
import os
import sys
from typing import Dict, List, Any, Optional, Union, Tuple, Callable
from functools import wraps, lru_cache
from collections import defaultdict, Counter
from dataclasses import dataclass, field
from enum import Enum, IntEnum
from contextlib import contextmanager

# Django core imports
from django.db import models, transaction
from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.cache import cache
from django.core.exceptions import ValidationError, PermissionDenied
from django.core.validators import MinValueValidator, MaxValueValidator, RegexValidator
from django.core.mail import send_mail
from django.core.files.uploadedfile import InMemoryUploadedFile
from django.core.serializers.json import DjangoJSONEncoder
from django.utils import timezone
from django.utils.translation import gettext_lazy as _
from django.utils.crypto import get_random_string, salted_hmac
from django.utils.encoding import force_bytes, force_str
from django.utils.functional import cached_property
from django.utils.module_loading import import_string
from django.urls import reverse, reverse_lazy
from django.http import HttpResponse, JsonResponse, HttpResponseRedirect
from django.shortcuts import render, get_object_or_404
from django.views.decorators.http import require_http_methods
from django.views.decorators.vary import vary_on_headers
from django.views.decorators.csrf import csrf_exempt, csrf_protect
from django.views.decorators.cache import cache_page, never_cache
from django.utils.decorators import method_decorator
from django.contrib import messages
from django.contrib.messages import constants as messages_constants
from django.db.models import Q, F, Count, Sum, Avg, Max, Min, Expression, Window, Func, Subquery, OuterRef, StdDev
from django.db.models.functions import Coalesce, Cast, Extract, Trunc, Now, TruncHour, TruncDay, ExtractHour
from django.db.models.signals import pre_save, post_save, pre_delete, post_delete, m2m_changed
from django.dispatch import receiver, Signal
from django.middleware.csrf import get_token
from django.template.loader import render_to_string
from django.forms import Form, ModelForm
from django.forms.fields import CharField, EmailField, ChoiceField, BooleanField
from django.forms.widgets import TextInput, PasswordInput, Select, CheckboxInput
from django.contrib.admin import ModelAdmin, site
from django.contrib import admin
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.sessions.models import Session
from django.contrib.contenttypes.models import ContentType
from django.contrib.contenttypes.fields import GenericForeignKey, GenericRelation
from django.contrib.postgres.fields import ArrayField, HStoreField
from django.db.models import JSONField
from django.contrib.postgres.indexes import GinIndex, BrinIndex
from django.contrib.postgres.search import SearchVector, SearchQuery, SearchRank
from django.contrib.postgres.aggregates import StringAgg

# Django REST Framework imports
from rest_framework import viewsets, generics, status, permissions, serializers, mixins
from rest_framework.decorators import action, api_view, permission_classes, renderer_classes, throttle_classes
from rest_framework.response import Response
from rest_framework.request import Request
from rest_framework.test import APITestCase, APIRequestFactory
from rest_framework.authtoken.models import Token
from rest_framework.authtoken.views import ObtainAuthToken
from rest_framework.pagination import PageNumberPagination, LimitOffsetPagination
from rest_framework.filters import SearchFilter, OrderingFilter, BaseFilterBackend
from rest_framework.permissions import BasePermission, IsAuthenticated, IsAdminUser
from rest_framework.serializers import Serializer, ModelSerializer, ValidationError
from rest_framework.fields import CharField as DRFCharField, EmailField as DRFEmailField
from rest_framework.validators import UniqueTogetherValidator
from rest_framework.exceptions import APIException, NotFound, PermissionDenied as DRFPermissionDenied
from rest_framework.settings import api_settings
from rest_framework.utils import formatting
from rest_framework.negotiation import DefaultContentNegotiation
from rest_framework.renderers import JSONRenderer, BrowsableAPIRenderer
from rest_framework.parsers import JSONParser, FormParser, MultiPartParser
from rest_framework.throttling import BaseThrottle, SimpleRateThrottle, AnonRateThrottle, UserRateThrottle
from rest_framework.views import APIView

# Celery imports
from celery import shared_task, Task
from celery.result import AsyncResult
from celery.schedules import crontab

# Cryptography imports
from cryptography.fernet import Fernet
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.backends import default_backend
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric import rsa, padding
from cryptography.hazmat.primitives.kdf.hkdf import HKDF
from cryptography.hazmat.primitives.hmac import HMAC

# Third-party imports
from faker import Faker
from factory import Factory, fuzzy, django, post_generation, Sequence, LazyAttribute
from pytest import fixture, mark
from unittest.mock import Mock, patch, MagicMock
from requests import Session, Request, Response
from requests.exceptions import RequestException, Timeout, ConnectionError
from PIL import Image, ImageDraw, ImageFont
from phonenumbers import parse, format_number, PhoneNumberFormat
from geoip2.database import Reader
from geoip2.errors import AddressNotFoundError, InvalidRequestError as InvalidDatabaseError
from urllib.parse import urlparse, urljoin, urlencode, parse_qs
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.header import Header
from email.utils import formataddr
from jwt import encode, decode, ExpiredSignatureError, InvalidTokenError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

# Local imports (conditional)
try:
    from .models import (
        DeviceInfo, SecurityLog, UserBan, IPBlacklist, ThreatIntelligence,
        SecurityIncident, TrustedDevice, PasswordHistory, SecurityPolicy,
        SecurityInsight, RiskScore, UserSession, TwoFactorMethod,
        SecurityNotification, AuditTrail
    )
    SECURITY_MODELS_AVAILABLE = True
except ImportError:
    SECURITY_MODELS_AVAILABLE = False

try:
    from .utils import SecurityUtils
    SECURITY_UTILS_AVAILABLE = True
except ImportError:
    SECURITY_UTILS_AVAILABLE = False

try:
    from .hash_chain import HashChainedModel
    HASH_CHAIN_AVAILABLE = True
except ImportError:
    HASH_CHAIN_AVAILABLE = False

try:
    from .permissions import BaseSecurityPermission
    PERMISSIONS_AVAILABLE = True
except ImportError:
    PERMISSIONS_AVAILABLE = False

# Common constants
DEFAULT_PAGE_SIZE = 20
MAX_PAGE_SIZE = 100
DEFAULT_CACHE_TIMEOUT = 300  # 5 minutes
SECURITY_CACHE_TIMEOUT = 3600  # 1 hour
RATE_LIMIT_PER_MINUTE = 60
RATE_LIMIT_PER_HOUR = 1000
RATE_LIMIT_PER_DAY = 10000

# Security constants
MIN_PASSWORD_LENGTH = 8
MAX_LOGIN_ATTEMPTS = 5
SESSION_TIMEOUT = 3600  # 1 hour
TOKEN_EXPIRATION = 86400  # 24 hours
HASH_SALT_LENGTH = 32
ENCRYPTION_KEY_LENGTH = 44  # Fernet key length

# Logging setup
security_logger = logging.getLogger('security')
audit_logger = logging.getLogger('security.audit')
auth_logger = logging.getLogger('security.auth')
threat_logger = logging.getLogger('security.threat')
performance_logger = logging.getLogger('security.performance')

# Common utility functions
def get_client_ip(request) -> str:
    """Get client IP address from request"""
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        ip = x_forwarded_for.split(',')[0]
    else:
        ip = request.META.get('REMOTE_ADDR')
    return ip or '0.0.0.0'

def generate_secure_token(length: int = 32) -> str:
    """Generate a cryptographically secure token"""
    return secrets.token_urlsafe(length)

def hash_string(data: str, salt: str = None) -> str:
    """Hash a string using SHA-256"""
    if salt:
        data = f"{data}{salt}"
    return hashlib.sha256(data.encode('utf-8')).hexdigest()

def verify_hash(data: str, hashed: str, salt: str = None) -> bool:
    """Verify a hash"""
    return hash_string(data, salt) == hashed

def is_valid_ip(ip: str) -> bool:
    """Check if IP address is valid"""
    try:
        ipaddress.ip_address(ip)
        return True
    except ValueError:
        return False

def is_private_ip(ip: str) -> bool:
    """Check if IP address is private"""
    try:
        ip_obj = ipaddress.ip_address(ip)
        return ip_obj.is_private
    except ValueError:
        return False

# Common decorators
def security_log(action: str, severity: str = 'info'):
    """Decorator for security logging"""
    def decorator(func):
        @wraps(func)
        def wrapper(*args, **kwargs):
            try:
                result = func(*args, **kwargs)
                security_logger.info(f"Security action: {action} - Success")
                return result
            except Exception as e:
                security_logger.error(f"Security action: {action} - Error: {str(e)}")
                raise
        return wrapper
    return decorator

def cache_result(timeout: int = DEFAULT_CACHE_TIMEOUT):
    """Decorator for caching results"""
    def decorator(func):
        @wraps(func)
        def wrapper(*args, **kwargs):
            cache_key = f"{func.__name__}:{hash(str(args) + str(kwargs))}"
            result = cache.get(cache_key)
            if result is None:
                result = func(*args, **kwargs)
                cache.set(cache_key, result, timeout)
            return result
        return wrapper
    return decorator

def rate_limit(limit: int, period: int = 60):
    """Decorator for rate limiting"""
    def decorator(func):
        @wraps(func)
        def wrapper(request, *args, **kwargs):
            client_ip = get_client_ip(request)
            cache_key = f"rate_limit:{func.__name__}:{client_ip}"
            count = cache.get(cache_key, 0)
            
            if count >= limit:
                from rest_framework.exceptions import Throttled
                raise Throttled(f"Rate limit exceeded: {limit} per {period} seconds")
            
            cache.set(cache_key, count + 1, period)
            return func(request, *args, **kwargs)
        return wrapper
    return decorator

# Common exceptions
class SecurityException(Exception):
    """Base security exception"""
    pass

class AuthenticationError(SecurityException):
    """Authentication error"""
    pass

class AuthorizationError(SecurityException):
    """Authorization error"""
    pass

class ValidationError(SecurityException):
    """Validation error"""
    pass

class RateLimitError(SecurityException):
    """Rate limit error"""
    pass

class SecurityConfigError(SecurityException):
    """Security configuration error"""
    pass

# Export all commonly used items
__all__ = [
    # Standard library
    'hashlib', 'logging', 'json', 'uuid', 'secrets', 'string', 're', 'ipaddress',
    'base64', 'hmac', 'time', 'datetime', 'decimal', 'os', 'sys',
    'Dict', 'List', 'Any', 'Optional', 'Union', 'Tuple', 'Callable',
    'wraps', 'lru_cache', 'defaultdict', 'Counter', 'dataclass', 'field',
    'Enum', 'IntEnum', 'contextmanager',
    
    # Django core
    'models', 'transaction', 'settings', 'get_user_model', 'Group', 'Permission',
    'cache', 'ValidationError', 'PermissionDenied', 'MinValueValidator', 'MaxValueValidator',
    'RegexValidator', 'send_mail', 'timezone', '_', 'get_random_string', 'salted_hmac',
    'force_bytes', 'force_str', 'cached_property', 'import_string', 'reverse', 'reverse_lazy',
    'HttpResponse', 'JsonResponse', 'HttpResponseRedirect', 'render', 'get_object_or_404',
    'require_http_methods', 'csrf_exempt', 'csrf_protect', 'cache_page', 'never_cache',
    'messages', 'messages_constants', 'Q', 'F', 'Count', 'Sum', 'Avg', 'Max', 'Min',
    'Expression', 'Window', 'Func', 'Coalesce', 'Cast', 'Extract', 'Trunc', 'Now',
    'pre_save', 'post_save', 'pre_delete', 'post_delete', 'm2m_changed', 'receiver', 'Signal',
    'get_token', 'render_to_string', 'Form', 'ModelForm', 'CharField', 'EmailField',
    'ChoiceField', 'BooleanField', 'TextInput', 'PasswordInput', 'Select', 'CheckboxInput',
    'ModelAdmin', 'site', 'admin', 'LoginRequiredMixin', 'PermissionRequiredMixin',
    'login_required', 'permission_required', 'Session', 'ContentType', 'GenericForeignKey',
    'GenericRelation', 'ArrayField', 'JSONField', 'HStoreField', 'GinIndex', 'BrinIndex',
    'SearchVector', 'SearchQuery', 'SearchRank', 'StringAgg',
    
    # Django REST Framework
    'viewsets', 'generics', 'status', 'permissions', 'serializers', 'mixins',
    'action', 'api_view', 'permission_classes', 'renderer_classes', 'Response', 'Request',
    'APITestCase', 'APIRequestFactory', 'Token', 'ObtainAuthToken', 'PageNumberPagination',
    'LimitOffsetPagination', 'SearchFilter', 'OrderingFilter', 'BaseFilterBackend',
    'BasePermission', 'IsAuthenticated', 'IsAdminUser', 'Serializer', 'ModelSerializer',
    'DRFCharField', 'DRFEmailField', 'UniqueTogetherValidator', 'APIException', 'NotFound',
    'DRFPermissionDenied', 'api_settings', 'formatting', 'DefaultContentNegotiation',
    'JSONRenderer', 'BrowsableAPIRenderer', 'JSONParser', 'FormParser', 'MultiPartParser',
    'APIView', 'BaseThrottle', 'SimpleRateThrottle', 'AnonRateThrottle', 'UserRateThrottle',
    
    # Celery
    'shared_task', 'Task', 'AsyncResult', 'crontab', 'task',
    
    # Cryptography
    'Fernet', 'hashes', 'PBKDF2HMAC', 'Cipher', 'algorithms', 'modes', 'default_backend',
    'serialization', 'rsa', 'padding', 'HKDF', 'HMAC',
    
    # Third-party
    'Faker', 'Factory', 'fuzzy', 'django', 'post_generation', 'Sequence', 'LazyAttribute',
    'fixture', 'mark', 'Mock', 'patch', 'MagicMock', 'Session', 'Request', 'Response',
    'RequestException', 'Timeout', 'ConnectionError', 'Image', 'ImageDraw', 'ImageFont',
    'parse', 'format_number', 'PhoneNumberFormat', 'Reader', 'AddressNotFoundError',
    'InvalidDatabaseError', 'urlparse', 'urljoin', 'urlencode', 'parse_qs', 'MIMEText',
    'MIMEMultipart', 'Header', 'formataddr', 'encode', 'decode', 'ExpiredSignatureError',
    'InvalidTokenError', 'Workbook', 'load_workbook', 'Font', 'PatternFill', 'Border',
    'Side', 'get_column_letter', 'canvas', 'letter', 'A4', 'inch', 'getSampleStyleSheet',
    'SimpleDocTemplate', 'Paragraph', 'Spacer',
    
    # Local imports (conditional)
    'SECURITY_MODELS_AVAILABLE', 'SECURITY_UTILS_AVAILABLE', 'HASH_CHAIN_AVAILABLE',
    'PERMISSIONS_AVAILABLE',
    
    # Constants
    'DEFAULT_PAGE_SIZE', 'MAX_PAGE_SIZE', 'DEFAULT_CACHE_TIMEOUT', 'SECURITY_CACHE_TIMEOUT',
    'RATE_LIMIT_PER_MINUTE', 'RATE_LIMIT_PER_HOUR', 'RATE_LIMIT_PER_DAY',
    'MIN_PASSWORD_LENGTH', 'MAX_LOGIN_ATTEMPTS', 'SESSION_TIMEOUT', 'TOKEN_EXPIRATION',
    'HASH_SALT_LENGTH', 'ENCRYPTION_KEY_LENGTH',
    
    # Logging
    'security_logger', 'audit_logger', 'auth_logger', 'threat_logger', 'performance_logger',
    
    # Utility functions
    'get_client_ip', 'generate_secure_token', 'hash_string', 'verify_hash', 'is_valid_ip',
    'is_private_ip',
    
    # Decorators
    'security_log', 'cache_result', 'rate_limit',
    
    # Exceptions
    'SecurityException', 'AuthenticationError', 'AuthorizationError', 'ValidationError',
    'RateLimitError', 'SecurityConfigError'
]
from rest_framework.throttling import UserRateThrottle, AnonRateThrottle, ScopedRateThrottle
# DeviceInfo stub — replace with real import if available
try:
    from api.security.models import DeviceInfo
except ImportError:
    class DeviceInfo:
        pass

